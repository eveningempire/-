from docx import Document
from pathlib import Path
p=Path('交付文档/重复使用运载器健康管理平台软件使用手册_202610版.docx'); d=Document(p)
text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
for key in ['\\nfigure','TODO','placeholder','未找到','截图文件未找到']:
 print(key,text.count(key))
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes),'chars',len(text))
