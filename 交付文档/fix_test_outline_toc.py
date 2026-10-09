from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm


src = Path(r'C:\Users\HP\Desktop\重复使用运载器健康管理原型系统测试大纲_合同指标补充版.docx')
out = Path(__file__).with_name('重复使用运载器健康管理原型系统测试大纲_合同指标补充版_目录修正.docx')
doc = Document(src)

# The source document restarts printed page numbers after the cover and TOC:
# PDF page 3 is printed page 2. These values were checked against the rendered
# source, rather than inferred from the PDF page count.
pages = {
    '1': 2, '2': 2, '3': 2, '4': 2,
    '5': 3, '5.1': 3, '5.2': 4,
    '6': 4, '6.1': 4, '6.2': 8,
    '6.2.1': 8, '6.2.2': 8,
    '6.2.2.1': 8, '6.2.2.2': 8,
    '6.2.2.3': 9, '6.2.2.4': 9,
    '6.2.3': 10, '6.2.4': 53,
    '6.2.5': 56, '6.2.6': 62,
    '7': 63,
}

toc_items = [
    ('1','适用范围'), ('2','测试依据'), ('3','测试组织'),
    ('4','测试范围'), ('5','测试环境'), ('5.1','软件架构'),
    ('5.2','硬件环境'), ('6','测试项目及标准'),
    ('6.1','齐套性检查'), ('6.2','软件类测试'),
    ('6.2.1','概述'), ('6.2.2','合同要求'),
    ('6.2.2.1','合同第四条研制内容与测试覆盖'),
    ('6.2.2.2','指标体系与方法数量'),
    ('6.2.2.3','健康管理原型系统设计'),
    ('6.2.2.4','原型系统性能'),
    ('6.2.3','功能测试'), ('6.2.4','性能测试'),
    ('6.2.5','可靠性和可维护性测试'),
    ('6.2.6','合格判定'), ('7','用例追踪关系'),
]

all_paras = list(doc.paragraphs)
toc_title = next(p for p in all_paras if p.text.replace('\u3000',' ').strip() == '目 录')
toc_old = [p for p in all_paras if p.style.name.startswith('toc ')]
assert len(toc_old) == 17, len(toc_old)
for p in toc_old:
    p._element.getparent().remove(p._element)

anchor = toc_title._p
for number, label in toc_items:
    level = min(number.count('.') + 1, 4)
    p = doc.add_paragraph(style=f'toc {level}')
    anchor.addnext(p._p)
    anchor = p._p
    # Match the original catalogue: number, tab, title, dot leader, right page.
    p.add_run(number + '\t' + label + '\t' + str(pages[number]))
    stops = p.paragraph_format.tab_stops
    # The original TOC uses separate number/title tabs. The first tab must
    # clear each style's left indent, especially for levels 3 and 4.
    number_tab = {1: 0.55, 2: 1.1, 3: 2.25, 4: 3.0}[level]
    stops.add_tab_stop(Cm(number_tab), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES)
    stops.add_tab_stop(Cm(14.7), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.paragraph_format.keep_together = True

doc.save(out)
print(out)
print('TOC entries:', len(toc_items))
