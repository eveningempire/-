from docx import Document
from pathlib import Path
p=Path('交付文档/重复使用运载器健康管理平台软件使用手册_202610版.docx')
d=Document(p)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes))
print('chars',sum(len(p.text) for p in d.paragraphs)+sum(len(c.text) for t in d.tables for row in t.rows for c in row.cells))
for i,p in enumerate(d.paragraphs):
 if p.text and ('TODO' in p.text or '未找到' in p.text or 'D:/PHM' in p.text): print(i,p.text)
