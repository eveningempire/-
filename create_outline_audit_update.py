# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(r"C:\Users\HP\Desktop\重复使用运载器健康管理原型系统测试大纲_平台适配审查版.docx")
doc = Document()
sec = doc.sections[0]
for attr in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(sec, attr, Cm(2.1))
styles = doc.styles
styles['Normal'].font.name='宋体'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); styles['Normal'].font.size=Pt(10.5)
for name,size in [('Title',18),('Heading 1',14),('Heading 2',12)]:
    styles[name].font.name='黑体'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); styles[name].font.size=Pt(size); styles[name].font.bold=True

def shade(cell, color='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def borders(table):
    tblPr=table._tbl.tblPr; b=tblPr.first_child_found_in('w:tblBorders')
    if b is None: b=OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=b.find(qn('w:'+edge))
        if e is None: e=OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'B7C3D0')
def table(title, ident, priority, purpose, pre, criterion, steps, note, status='已按当前实现修订'):
    doc.add_heading(title, level=1)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f'测试用例标识：{ident}').bold=True
    info=doc.add_table(rows=0,cols=2); info.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(info)
    fields=[('测试用例名称',title),('测试用例标识',ident),('测试优先级',priority),('被测软件名称','健康管理原型系统软件'),('被测软件版本标识','现场交付版本（测试时填写）'),('测试目的',purpose),('测试类型','☑功能测试  □性能测试  □接口测试  □安全性测试  □保密性测试  □适应性测试  □文档审查'),('测试方法','□白盒测试  ☑黑盒测试'),('前提和约束（包括初始化要求）',pre),('测试终止条件','完成全部步骤，或因异常无法继续时终止；终止时记录原因及已完成步骤。'),('测试用例通过准则',criterion)]
    for k,v in fields:
        c=info.add_row().cells; c[0].text=k; c[1].text=v; shade(c[0]); c[0].vertical_alignment=c[1].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_heading('测试规程',level=2)
    t=doc.add_table(rows=1,cols=5); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for c,h in zip(t.rows[0].cells,['序号','输入及操作步骤','预期测试结果','实际测试结果','判定']): c.text=h; shade(c,'5B9BD5'); c.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255); c.paragraphs[0].runs[0].bold=True
    for n,a,e in steps:
        cells=t.add_row().cells
        for c,v in zip(cells,[str(n),a,e,'','']): c.text=v; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for k,v in [('问题标识',''),('测试结论','□通过      □未通过      □未测试'),('测试人员','                测试时间：'),('备注',note)]:
        c=t.add_row().cells; c[0].text=k; c[1].merge(c[4]); c[1].text=v; shade(c[0])
    doc.add_paragraph()

doc.add_heading('重复使用运载器健康管理原型系统测试大纲平台适配审查版',0)
doc.add_paragraph('本版依据当前工作区平台代码、实时接口、已登记数据和 2026 年 10 月 9 日真实仿真验证记录编写。平台代码未修改。本版的“通过”仅表示对应功能和证据满足该用例；没有独立真值、现场性能测量或正式交付材料时，必须记录为“未测试”或“未通过”，不能以页面可见或接口返回代替验收。')
doc.add_heading('一、平台核对结论',1)
for text in [
'已具备并可按步骤执行：数据导入、D01/D02 历史回放、回放控制、Simulink 任务提交与 CSV 登记、实时采集会话、默认告警规则、PBS 结构管理、批量数据导入、模型导入、权限控制、三维页面和多信号流图。',
'需要现场证据后才能判定：启动 ≤60 秒、接收至显示最大延迟 ≤1 秒、完整审计字段、交付文档齐套、故障树至少 1 例、指标体系及逐级分配、健康评估标定精度、寿命预测精度。',
'当前不应直接判定通过：MSFG+TEAMS-RT 和 PCA-iForest 的一院故障分类正确性、所有故障的 Hier14 正确性、没有真实 RUL 标签时的寿命精度。真实验证记录显示这些项目存在误报、漏报或缺少真值。',
'已知口径限制：Simulink 故障接口使用故障类型、注入时刻、恢复时刻和飞行次数，不使用严重度；仿真完成后一次性导出 CSV，不提供逐点在线求解遥测；寿命预测训练时间单位为秒；结构树与算法归属通过显式 PBS 关联，不能仅靠系统或部件文字推断。']:
    doc.add_paragraph(text,style='List Bullet')

table('6.2.3.4 故障注入与清除恢复测试用例','JKGLYXXT-GN-02-01','重要','验证真实 Simulink 故障仿真任务的参数校验、后台求解、阶段反馈、CSV 登记及任务结果关联。','测试计算机及 MATLAB/Simulink 运行环境正常，许可证可用；平台 Engine 状态可用；用户具有 submit_simulation 权限。准备一个动力故障和一个控制故障；注入时刻和恢复时刻满足 0 ≤ 注入时刻 < 恢复时刻 ≤ 480 s。当前接口不设置严重度字段。','动力和控制场景完成任务创建、后台执行、阶段更新、原始 CSV 登记和结果查看；FaultFlag 区间与设定时间一致；失败任务有错误并可重试。不要求逐点在线 MATLAB 遥测。',[(1,'选择动力故障，设置注入 10 s、恢复 20 s、飞行次数 0，提交仿真。','生成任务 UUID，进入排队/运行状态并显示启动 Engine、配置模型、求解、导出、登记阶段。'),(2,'轮询至完成，查看仿真输出。','显示原始状态量 CSV、数据集编号、样本数、FaultFlag 数量、参数数量和下载链接。'),(3,'预览/下载 CSV，核对 Time、FaultFlag 和信号；在数据管理或故障回放打开。','Time 为秒且递增；FaultFlag 按区间规则记录；结果属于当前任务。'),(4,'新建控制故障任务，设置注入 30 s、恢复 40 s、飞行次数 0，完成并查看。','控制场景独立完成；两个任务的 UUID、参数、数据集和结果不串用。'),(5,'提交恢复≤注入或超出 0～480 s 的任务；对失败任务重试。','非法参数明确拒绝；失败原因可见；修正后可重试，已完成任务不重复启动。')],'没有单独“清除故障并恢复”按钮；恢复通过 end_time 结束故障区间实现。仿真完成后一次性导出 CSV。隔离提示不表示实体设备隔离。')

table('6.2.3.5 告警阈值与恢复判定测试用例','JKGLYXXT-GN-02-02','重要','验证实时采集会话的启用控制、默认告警规则边界判定、告警关联和历史可追溯性。','创建独立 MATLAB 实时采集会话并保存一次性令牌；在异常告警中心选择该会话并启用实时检测；time 严格递增且字段集合不变。默认 pressure < 0.88 触发“推进剂泄漏”严重告警；检测不依赖 FaultFlag。','关闭检测不上告警；开启后正常、等号、越限和恢复按规则判定。告警包含会话、采样、信号、值、阈值、操作符、等级、故障名称、隔离对象和健康指数。',[(1,'检测关闭时上报 time=0、pressure=0.70、temperature=0.40。','接收成功但不产生告警。'),(2,'开启检测，递增 time 上报 pressure=1.00、0.88。','均不触发 pressure<0.88。'),(3,'上报 pressure=0.879、0.70，查询告警中心及接口。','每个越限采样产生告警；阈值、操作符、等级、故障名称、隔离对象和健康指数正确。'),(4,'上报 pressure=1.00，查询当前与历史告警。','恢复点不新增同类越限告警；原告警仍可追溯。'),(5,'用错误令牌、重复/倒退 time、改变字段集合请求。','分别返回 403 或明确 400 错误且不写入。')],'默认规则还包括 temperature > 0.52、attitude_error > 0.08、control_error > 0.10。每个命中的采样生成告警，不自动合并相邻越限点。')

doc.add_heading('二、其他必须在现场修正判定口径的用例',1)
rows=[('FF-FTA-01','故障树至少有 1 例完整记录；当前数据库未发现已保存故障树，只有编辑接口，因此准备故障树并保存后才能测试。'),('FF-DIAG-01～04','先判接口执行和结果展示，再用独立真值样本判准确率；当前真实记录不能支持全部算法通过。'),('FF-LIFE-01','页面和接口按秒（s）记录；大纲不得要求显示小时，除非提供经确认的换算依据。没有真实 RUL 标签只能判输入/执行兼容，不能判精度。'),('FF-HEALTH-01～03','CDPCA-GA、AE-GMM、GCN/RBD 可执行，但目前缺少型号标定健康真值；精度项记未测试。'),('FF-EXT-01','批量导入最多 2000、诊断最多 200、寿命最多 100；属于客户端分批同步处理，不能写成每个文件都有独立后台任务，现场应按容量限制判定。'),('JKGLYXXT-GN-08-01','三维页面包含预置仿真单元信息；若要求与实际仿真任务联动，需现场建立任务并核对，静态预置内容只能判展示通过。'),('JKGLYXXT-XN-01/02','当前没有三轮启动计时和 6000 点最大延迟测量记录；必须现场按大纲工具和误差上界测量，否则记未测试。'),('ZL-05','当前有 AuditEvent 模型及部分遥测写入记录，但不应据此推断所有数据、模型、权限和失败操作均完整审计；需逐项核对操作者、时间、对象和结果。'),('QT-03/QT-04','平台代码和 Markdown 说明不能替代正式签署的需求、设计、研究总结、报价/合同成果；按交付包逐项核对并记录缺件。')]
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t); t.rows[0].cells[0].text='用例'; t.rows[0].cells[1].text='适配结论'; shade(t.rows[0].cells[0],'5B9BD5'); shade(t.rows[0].cells[1],'5B9BD5')
for a,b in rows: c=t.add_row().cells; c[0].text=a; c[1].text=b
doc.add_heading('三、测试记录要求',1)
doc.add_paragraph('每项测试记录版本号、机器配置、账号角色、输入文件哈希、算法/模型版本、请求与响应、截图、日志、实际结果和判定。功能可打开不等于指标通过；缺少参考真值、标定数据、现场计时或正式交付件时，结论选择“未测试”，并填写补充条件。')
doc.save(str(out)); print(out)
