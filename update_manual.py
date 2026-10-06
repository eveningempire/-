from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from shutil import copyfile

src = r'C:\Users\HP\Desktop\重复使用运载器健康管理平台软件使用手册 (2).docx'
out = r'E:\实验项目\航天管理项目\phm_platform\重复使用运载器健康管理平台软件使用手册_离线交付版.docx'
copyfile(src, out)
doc = Document(out)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def add_heading(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Microsoft YaHei'
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body(text, bold_lead=None):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def add_code(text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    return p

doc.add_page_break()
add_heading('20 离线交付与部署', 1)
add_body('本章适用于无互联网条件下部署本平台的甲方计算机。离线交付包已经包含 Python 3.10.11 安装器、Node.js 安装器、项目虚拟环境、前端构建产物、项目源代码和启动脚本。现场部署不需要访问 PyPI、npm 或其他外部网络。')

add_heading('20.1 交付包内容', 2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr = table.rows[0].cells
for i, t in enumerate(['目录或文件', '用途', '是否现场必需']):
    hdr[i].text = t
    set_cell_shading(hdr[i], 'D9EAF7')
for row in [
    ('runtime-installers\\python-3.10.11-amd64.exe', '安装 Python 3.10.11', '是'),
    ('runtime-installers\\node-v22.18.0-x64.msi', '前端重建或迁移模式使用', '按需'),
    ('.venv', '已经验证的 Python 运行环境', '是'),
    ('frontend\\dist', '已构建的前端页面', '是'),
    ('frontend\\node_modules', '前端离线依赖', '按需'),
    ('db.sqlite3', 'SQLite 业务数据库', '按需'),
    ('install_runtime_offline.bat、setup_migration.ps1、start_cmg_platform.bat', '安装、初始化和启动脚本', '是'),
]:
    cells = table.add_row().cells
    for i, t in enumerate(row):
        cells[i].text = t

add_heading('20.2 现场部署步骤', 2)
add_body('1. 将 phm_platform_offline_delivery.zip 复制到甲方计算机并解压到没有中文或空格的目录,例如 D:\\PHM\\phm_platform。')
add_body('2. 进入解压后的项目根目录,确认能够看到 manage.py、.venv、runtime-installers 和 frontend 目录。')
add_body('3. 双击 install_runtime_offline.bat。该脚本从包内安装 Python 3.10.11 和 Node.js 22.18.0,不访问互联网。安装完成后关闭当前命令行窗口,并重新打开一个 PowerShell 窗口。')
add_body('4. 在项目根目录运行 setup_migration.ps1。脚本会优先检查随包提供的 .venv;如果环境完整,将跳过网络安装依赖,执行数据库和目录初始化。')
add_code('powershell -ExecutionPolicy Bypass -File .\\setup_migration.ps1')
add_body('5. 初始化管理员账号。首次部署且没有管理员账号时运行:')
add_code('.\\.venv\\Scripts\\python.exe create_superuser.py')
add_body('6. 双击 start_cmg_platform.bat,等待浏览器打开,或访问 http://127.0.0.1:8000/。')

add_heading('20.3 离线运行方式', 2)
add_body('日常使用只需要运行 start_cmg_platform.bat。脚本会使用项目内的 .venv 和 frontend\\dist,不会执行 pip install、npm install 或联网检查。平台默认使用 SQLite 数据库,Redis 和前端开发服务器不是普通使用的必要条件。')
add_body('停止平台时运行 stop_cmg_platform.bat,或关闭启动脚本打开的服务窗口。停止服务前应先完成当前数据采集、文件上传和数据库写入。')

add_heading('20.4 局域网访问和 MATLAB 采集', 2)
add_body('如果 MATLAB 或其他采集设备与平台部署在同一局域网,管理员可以将后端绑定到 0.0.0.0,并在 Windows 防火墙中放行 8000 端口。访问地址应使用部署计算机的实际局域网 IP,例如 http://192.168.1.20:8000/。')
add_code('.\\.venv\\Scripts\\python.exe manage.py runserver 0.0.0.0:8000')
add_body('甲方网络策略仍然可以保持完全禁止互联网访问;上述设置只用于内部局域网通信。')

add_heading('20.5 离线故障处理', 2)
for text in [
    '提示找不到 Python:确认已运行 install_runtime_offline.bat,并重新打开命令行;检查 runtime-installers\\python-3.10.11-amd64.exe 是否存在。',
    '提示 .venv 不完整:确认压缩包完整解压,不能只复制部分目录;重新解压后再运行 setup_migration.ps1。',
    '浏览器打不开 8000 端口:检查 start_cmg_platform.bat 是否仍在运行,确认端口没有被其他程序占用,并检查 Windows 防火墙规则。',
    '页面显示异常:优先使用 start_cmg_platform.bat 启动已打包的 frontend\\dist;只有修改了前端源码时,才需要安装 Node.js 后执行 npm run build。',
    '需要重建 Python 环境:在能够联网的制作机上补齐 offline\\python 下的全部 wheel 文件,再在现场使用 pip --no-index --find-links offline\\python 安装。现场不应临时访问互联网。',
]:
    add_body(text, text.split(':')[0] + ':')

add_heading('20.6 交付验收记录', 2)
for text in [
    '□ 甲方计算机未连接互联网时,Python 3.10.11 安装完成。',
    '□ setup_migration.ps1 能够完成初始化并保持命令行无联网错误。',
    '□ start_cmg_platform.bat 能够启动平台并打开登录页面。',
    '□ 管理员账号能够登录,综合态势页面能够加载。',
    '□ 数据库、文件上传和核心业务页面能够正常使用。',
]:
    add_body(text)

doc.save(out)
print(out)
