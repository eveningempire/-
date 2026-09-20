from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE = Path(r"E:\实验项目\航天管理项目\交付包\项目使用说明.md")
OUTPUT = Path(r"E:\实验项目\航天管理项目\交付包\项目使用说明.docx")

BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
GRID = "D9D9D9"


def set_font(run, latin="Arial", east_asia="Microsoft YaHei", size=None, bold=None, color=BLACK):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), GRID)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color="666666")
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_sep, text, fld_end])
    tail = paragraph.add_run(" 页")
    set_font(tail, size=9, color="666666")


def add_inline_text(paragraph, text, size=10.5):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, latin="Consolas", east_asia="Microsoft YaHei", size=9.5, color="7A1F1F")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(16)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 11.5, 10, 5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string("202020")
    code.paragraph_format.left_indent = Cm(0.45)
    code.paragraph_format.right_indent = Cm(0.3)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.1


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("重复使用运载器健康管理平台")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("客户使用说明")
    set_font(r, size=19, bold=True)
    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Cm(2)
    intro.paragraph_format.right_indent = Cm(2)
    rr = intro.add_run("本手册用于指导客户完成平台安装、启动、日常操作、MATLAB 实时数据接入、备份及常见问题处理。")
    set_font(rr, size=11, color="404040")
    for _ in range(6):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.line_spacing = 1.7
    for line in ("文档版本  1.0", "更新日期  2026 年 9 月 18 日", "适用系统  Windows 10 和 Windows 11 64 位"):
        run = meta.add_run(line + "\n")
        set_font(run, size=10.5, color="4D4D4D")
    doc.add_page_break()


def add_toc(doc, headings):
    p = doc.add_paragraph("目录", style="Heading 1")
    p.paragraph_format.space_after = Pt(14)
    for level, text in headings:
        if level != 1:
            continue
        row = doc.add_paragraph()
        row.paragraph_format.left_indent = Cm(0)
        row.paragraph_format.space_after = Pt(4)
        run = row.add_run(text)
        set_font(run, size=11, bold=True, color="202020")
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    r = note.add_run("提示：在 Word 中可通过导航窗格按标题快速定位章节。")
    set_font(r, size=9, color="666666")
    doc.add_page_break()


def parse_source(lines):
    headings = []
    for line in lines:
        m = re.match(r"^(#{2,3})\s+(.+)$", line)
        if m:
            headings.append((len(m.group(1)) - 1, m.group(2).strip()))
    return headings


def add_table(doc, raw_rows):
    rows = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in raw_rows]
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
        rows.pop(1)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(5.0), Cm(10.5)] if len(rows[0]) == 2 else [Cm(15.5 / len(rows[0]))] * len(rows[0])
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_text(p, value, size=9.2)
            if i == 0:
                set_cell_shading(cell, NAVY)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            elif i % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
    set_table_borders(table)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(doc, lines):
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                p.paragraph_format.keep_together = True
                set_cell = OxmlElement("w:shd")
                set_cell.set(qn("w:fill"), PALE_GRAY)
                p._p.get_or_add_pPr().append(set_cell)
                r = p.add_run("\n".join(code_lines))
                set_font(r, latin="Consolas", east_asia="Microsoft YaHei", size=8.5)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        m = re.match(r"^(#{1,3})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            if level == 1:
                i += 1
                continue
            text = m.group(2).strip().replace("——", " ")
            doc.add_paragraph(text, style=f"Heading {level - 1}")
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline_text(p, bullet.group(1))
        elif numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline_text(p, numbered.group(2))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            add_inline_text(p, line)
        i += 1


def add_headers_footers(doc):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("重复使用运载器健康管理平台  客户使用说明")
        set_font(r, size=8.5, color="666666")
        add_page_field(section.footer.paragraphs[0])


def main():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    headings = parse_source(lines)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_toc(doc, headings)
    add_body(doc, lines)
    add_headers_footers(doc)
    doc.core_properties.title = "重复使用运载器健康管理平台客户使用说明"
    doc.core_properties.subject = "平台安装 操作 MATLAB 接入 备份与故障排查"
    doc.core_properties.author = "PHM 平台项目组"
    doc.core_properties.keywords = "PHM MATLAB Simulink 客户手册"
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
