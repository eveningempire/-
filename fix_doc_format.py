from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING
from pathlib import Path

src=Path(r'C:\Users\HP\Desktop\重复使用运载器健康管理原型系统测试大纲_MATLAB与实时告警补充版.docx')
out=Path(r'C:\Users\HP\Desktop\重复使用运载器健康管理原型系统测试大纲_MATLAB与实时告警补充版_格式修正版.docx')
doc=Document(str(src))

def set_row_cant_split(row):
    trPr=row._tr.get_or_add_trPr()
    for old in trPr.findall(qn('w:cantSplit')):
        trPr.remove(old)
    trPr.append(OxmlElement('w:cantSplit'))

def set_repeat(row):
    trPr=row._tr.get_or_add_trPr()
    if trPr.find(qn('w:tblHeader')) is None:
        trPr.append(OxmlElement('w:tblHeader'))

def set_cell_padding(cell, top=70, start=90, bottom=70, end=90):
    tcPr=cell._tc.get_or_add_tcPr()
    tcMar=tcPr.find(qn('w:tcMar'))
    if tcMar is None:
        tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); tcMar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')

def set_cell_text(cell):
    for p in cell.paragraphs:
        pf=p.paragraph_format
        pf.space_before=Pt(0); pf.space_after=Pt(0)
        pf.line_spacing=1.08
        for run in p.runs:
            run.font.size=Pt(9)
            run.font.name='宋体'
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')

for table in doc.tables:
    # Preserve the source column structure, but ensure rows are not split at page boundaries.
    for row in table.rows:
        set_row_cant_split(row)
        for cell in row.cells:
            set_cell_padding(cell)
            set_cell_text(cell)
    for row in table.rows:
        if any('序号' in (cell.text or '') for cell in row.cells):
            set_repeat(row)
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold=True
                        run.font.size=Pt(9)

# Keep headings and table captions with the following table where possible.
for p in doc.paragraphs:
    if p.text.strip().startswith(('表 ', '表\u00a0')) or '测试用例' in p.text:
        p.paragraph_format.keep_with_next=True

doc.save(str(out))
print(out)
