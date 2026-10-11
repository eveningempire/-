# -*- coding: utf-8 -*-
from copy import deepcopy
from pathlib import Path
from docx import Document

SRC = Path(r"C:\Users\HP\Desktop\重复使用运载器健康管理原型系统测试大纲.docx")
OUT = Path(r"C:\Users\HP\Desktop\重复使用运载器健康管理原型系统测试大纲_平台适配完整修订版.docx")

def clean(s):
    return ' '.join((s or '').split())

def physical_cells(row):
    # One XML cell may be represented by several python-docx cells when a row
    # contains horizontal merges; use the first instance of each tc element.
    result=[]; seen=set()
    for cell in row.cells:
        key=id(cell._tc)
        if key not in seen:
            seen.add(key); result.append(cell)
    return result

def field_cell(table, label):
    for row in table.rows:
        cells=physical_cells(row)
        if not cells: continue
        if label in clean(cells[0].text):
            return cells[-1]
    return None

def set_field(table, label, value):
    cell=field_cell(table,label)
    if cell is None: raise RuntimeError(f'missing field {label}')
    cell.text=value

def table_for(doc, ident):
    for t in doc.tables:
        if ident in ' '.join(c.text for row in t.rows for c in row.cells):
            return t
    raise RuntimeError(ident)

def set_steps(table, values):
    header_i=None
    for i,row in enumerate(table.rows):
        if any('输入及操作步骤' in clean(c.text) for c in physical_cells(row)):
            header_i=i; break
    if header_i is None: raise RuntimeError('step header missing')
    step_rows=[]
    for row in table.rows[header_i+1:]:
        cells=physical_cells(row)
        if cells and any(x in clean(cells[0].text) for x in ('问题标识','测试结论','测试人员','备 注','备注')):
            break
        if len(cells)>=5: step_rows.append(row)
    for row, vals in zip(step_rows, values):
        cells=physical_cells(row)
        for cell, value in zip(cells[:5], vals): cell.text=str(value)

def set_note(table, value):
    set_field(table, '备', value)

doc=Document(str(SRC))

# The two cases most affected by the changed fault-injection implementation.
t=table_for(doc,'JKGLYXXT-GN-02-01')
set_field(t,'测试目的','验证真实 Simulink 故障仿真任务的参数校验、后台求解、阶段反馈、CSV 登记及任务结果关联。')
set_field(t,'前提和约束','硬件约束条件：测试计算机及 MATLAB/Simulink 运行环境正常，许可证可用。软件约束条件：平台 Engine 状态可用，用户具有 submit_simulation 权限。数据及初始化要求：管理员登录，打开“故障注入”页面，准备一个动力故障和一个控制故障；注入时刻和恢复时刻单位为秒，满足 0 ≤ 注入时刻 < 恢复时刻 ≤ 480；当前接口不设置严重度字段。')
set_field(t,'测试用例通过准则','动力和控制场景均完成任务创建、后台执行、阶段状态更新、原始状态量 CSV 登记和结果查看；FaultFlag 区间与设定的注入/恢复时刻一致；失败任务保存可定位错误并可重试。不要求逐点在线 MATLAB 遥测，也不以任务完成代替诊断正确性。')
set_steps(t,[
('1','进入“故障注入”页面，刷新运行环境；选择动力故障，设置注入时刻 10 s、恢复时刻 20 s、飞行次数 0，点击“提交仿真”。','生成任务 UUID，任务进入排队或运行状态，并显示 MATLAB 启动、模型配置、求解、导出和登记阶段。','',''),
('2','轮询任务状态至完成，查看任务摘要和“仿真输出”。','任务进入“已完成”；输出包含原始状态量 CSV、数据集编号、样本数、故障标记样本数、参数数量和下载链接。','',''),
('3','预览或下载 CSV，核对 Time、FaultFlag 和信号；在数据管理或故障回放中打开该数据集。','Time 为秒且单调递增；FaultFlag 按当前模型的注入区间规则记录；数据集可追溯到当前任务。','',''),
('4','新建控制故障任务，设置注入 30 s、恢复 40 s、飞行次数 0，提交并等待完成；另提交恢复时刻不晚于注入时刻的非法任务。','控制场景独立完成，两个任务的 UUID、参数、数据集和结果不串用；非法参数明确拒绝，失败任务保留错误并可重试。','','')])
set_note(t,'真实接口使用故障类型、注入时刻、恢复时刻和飞行次数，不使用旧版“严重度”参数，也没有单独的“清除故障并恢复”按钮；恢复通过 end_time 结束故障区间实现。仿真完成后一次性导出 CSV，不提供逐点在线求解遥测。隔离提示不表示实体设备隔离。')

t=table_for(doc,'JKGLYXXT-GN-02-02')
set_field(t,'测试目的','验证实时采集会话的启用控制、默认告警规则边界判定、告警记录关联和恢复后的历史可追溯性。')
set_field(t,'前提和约束','硬件约束条件：测试计算机及平台服务正常。软件约束条件：用户具有创建采集会话和查看告警的权限。数据及初始化要求：管理员创建独立 MATLAB 实时采集会话并保存一次性令牌；在“异常告警中心”选择该会话并启用实时检测；time 严格递增且字段集合不变。当前默认规则为 pressure < 0.88 触发“推进剂泄漏”严重告警；检测不依赖 FaultFlag。')
set_field(t,'测试用例通过准则','实时检测关闭时不上报告警；开启后正常点、阈值等号点、越限点和恢复点按当前规则判定。告警包含会话、采样、信号、值、阈值、操作符、等级、故障名称、隔离对象和健康指数；历史告警可查询。')
set_steps(t,[
('1','检测关闭时上报 time=0、pressure=0.70、temperature=0.40；开启检测后按递增 time 上报 pressure=1.00、0.88。','关闭时数据接收成功但不产生告警；开启后 1.00 和等号 0.88 均不触发 pressure<0.88。','',''),
('2','继续上报 pressure=0.879、0.70，查询告警中心及会话告警接口。','每个越限采样产生告警；阈值、操作符、等级、故障名称、隔离对象与会话一致；健康指数为 0～1 有限值。','',''),
('3','上报恢复点 pressure=1.00；再使用错误令牌、重复或倒退 time、改变字段集合的请求进行边界验证。','恢复点不新增同类越限告警，原告警仍可追溯；错误令牌返回 403，非法时间或字段返回明确 400 错误且不写入。','','')])
set_note(t,'默认规则还包括 temperature > 0.52、attitude_error > 0.08、control_error > 0.10。自定义规则可能与默认规则叠加，应记录启用规则。系统按每个命中的采样生成告警，不自动合并相邻越限点。')

# Cases where the original wording overstated current evidence or used a wrong unit.
updates={
'FF-FTA-01': {
 'pre':'准备至少2类典型系统的正式故障分析资料、结构映射和至少1例待测故障树；故障树必须已保存或在测试前完成保存。当前平台虽有故障树编辑接口，但没有可直接作为验收证据的预置故障树。',
 'criterion':'至少2类典型系统、至少1例完整故障树；顶事件、逻辑门、中间事件、底事件、关联系统及推理路径均可复核。缺少已保存故障树时记未测试，不以打开编辑页面代替通过。',
 'note':'故障树数量和内容必须以保存记录及导出证据为准。'
},
'FF-DIAG-01': {'criterion':'先确认 Hier14 输入字段、模型版本和窗口规则，再分别核对执行结果与独立真值；仅接口返回成功不构成准确率通过。','note':'当前真实联调中部分故障类别和正常样本存在误报/漏报，需按独立测试集逐项记录，不得外推为全部故障通过。'},
'FF-DIAG-02': {'criterion':'MSFG+TEAMS-RT 能执行并输出证据，且在测试前确认的测点映射、阈值和独立真值下满足判据；无映射或仅有通用异常分数时记未测试。','note':'当前实现未完成一院13类故障的测点映射和工况标定，真实代表集出现正常基线误报；不能直接判准确性通过。'},
'FF-DIAG-03': {'criterion':'PCA-iForest 的输入、阈值和窗口结果可复核，并按独立真值样本判定；算法可运行不等于故障检出率合格。','note':'当前验证中正常窗口误报和泵气蚀检出不足，需补充独立正常/故障样本及标定阈值。'},
'FF-DIAG-04': {'criterion':'严重故障样本的识别、隔离、触发时间和恢复均与独立真值一致；漏检、误判和延迟必须记录为问题，不能用 FaultFlag 直接作为算法输出。','note':'本项要求现场准备严重故障真值和适用算法，当前资料不足时记未测试。'},
'FF-LIFE-01': {'pre':'准备符合模型定义的 HI 或可由匹配仿真派生参数生成的输入、至少51行有效序列、相应 PBS 部件模型、参考 RUL 及容差。','criterion':'至少1类系统、1种方法可执行，输入配对、部件映射、HI 预处理和单位均符合模型契约；没有真实 RUL 标签时只判输入/执行兼容，不判预测精度。','note':'当前训练模型返回单位为秒（s），不是小时（h）；除非取得经确认的换算依据，不得在测试记录中改写为小时。新数据存在严重 HI 裁剪时，结果不能作为标定寿命结论。'},
'FF-HEALTH-01': {'criterion':'CDPCA-GA 可执行并输出本次 HI 序列、曲线和坐标；有独立参考 HI 与容差时再判精度，否则记输入/执行通过、精度未测试。','note':'当前健康指数属于相对评估，缺少型号/工况标定真值。'},
'FF-HEALTH-02': {'criterion':'AE-GMM 独立执行并输出 HI 和状态；有独立参考输出时按容差判精度，否则记未测试。','note':'需记录训练迭代、模型资产和异常输入；能返回结果不证明健康指数准确。'},
'FF-HEALTH-03': {'criterion':'GCN/RBD 能按显式 PBS 关系融合至少两个有效部件并输出系统 HI；有参考系统级真值时判精度，否则只判接口和结构一致性。','note':'原始/派生文件属于同一次物理运行，不能作为两个独立部件样本。'},
'FF-EXT-01': {'criterion':'在约定容量内完成批量导入/诊断/寿命处理；成功、失败和限制提示总数一致，点击条目显示对应结果。当前接口为同步分批处理，不要求每个文件都有独立后台任务。','note':'单次批量导入上限2000、批量诊断上限200、批量寿命预测上限100；超出容量应由客户端分批提交。'},
'JKGLYXXT-GN-08-01': {'criterion':'静态预置单元可展示名称、状态和反馈且明确标注演示数据；若要求与实际仿真任务联动，必须现场提交任务并核对变化，预置内容不能代替联动证据。','note':'当前三维页面的仿真单元来自前端预置配置，不能仅凭页面显示证明与后台 Simulink 任务联动。'},
'JKGLYXXT-XN-01': {'criterion':'三轮从启动脚本开始至进入工作状态的实测时间均≤60秒；无三轮计时记录时记未测试。','note':'依赖安装、迁移和前端构建应在计时前完成，但启动脚本实际初始化仍计入。'},
'JKGLYXXT-XN-02': {'criterion':'在1个采集源、1个查看端、6个信号、10点/秒、10分钟、6000点工况下，最大显示延迟加测量误差上界≤1秒；无自动记录和误差校准证据时记未测试。','note':'页面最近1000点限制不改变全程采集记录；应保存6000点原始记录和接收/渲染时间。'},
'ZL-05': {'criterion':'关键操作实际产生的审计记录可追溯到操作者、时间、对象和结果；若当前版本只记录部分遥测或服务器日志，未覆盖的操作记未测试/未通过，不以接口访问日志代替审计。','note':'需逐项核对登录、数据、模型、权限、任务成功/失败和告警操作；日志不得泄露密码或完整 Token。'},
'JKGLYXXT-QT-03': {'criterion':'正式需求规格说明报告有版本和签署页，覆盖合同、纪要和指标并能追踪到用例；缺正式签署件或缺条款时未通过。','note':'平台 Markdown 和页面说明不能替代正式交付报告。'},
'JKGLYXXT-QT-04': {'criterion':'正式软件设计方案有版本和签署页，架构、接口、扩展点、批量限制和异常路径可定位并与实装一致；缺正式签署件或内容不一致时未通过。','note':'需同时核对需求、设计、合同/报价及交付清单。'}
}
for ident, spec in updates.items():
    t=table_for(doc,ident)
    print('updating',ident)
    for key in ('pre','criterion','note'):
        if key not in spec: continue
        if key=='pre': set_field(t,'前提和约束',spec[key])
        elif key=='criterion': set_field(t,'测试用例通过准则',spec[key])
        else: set_note(t,spec[key])

# Append a concise auditable conclusion without changing any platform code.
doc.add_page_break()
doc.add_heading('平台适配审查结论',0)
doc.add_paragraph('本版只修订测试口径和证据要求，不代表平台已经通过全部验收项目。现场执行时，功能可操作、接口可返回与指标/精度/交付物达标分别记录。')
for item in [
'可直接执行的功能类：D01/D02 历史回放、回放控制、数据导入、Simulink 任务提交与 CSV 登记、实时采集与默认告警、PBS 管理、模型导入和权限接口。',
'必须补证据后判定：故障树、指标体系及逐级分配、健康评估标定、寿命预测精度、启动和刷新延迟、完整日志审计、正式需求/设计/合同成果。',
'当前不应预判通过：所有算法诊断正确性、严重故障识别、没有真实 RUL 标签时的寿命精度。',
'记录规范：每项保存版本、机器环境、账号角色、输入文件哈希、模型版本、请求响应、截图、日志、实际结果和判定；缺参考真值、标定数据、现场测量或正式文件时选择“未测试”并说明补充条件。']:
    doc.add_paragraph(item, style='List Bullet')
doc.save(str(OUT))
print(OUT)
