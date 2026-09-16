# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r'E:\实验项目\航天管理项目\phm_platform\docs\健康管理平台验收操作与答辩话术手册.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
for name,size in [('Normal',10.5),('Title',22),('Heading 1',16),('Heading 2',13)]:
    s=doc.styles[name]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
def shade(c,fill):
    tcPr=c._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def tbl(headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E79')
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for n,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=str(v)
            if n%2: shade(cs[i],'F3F6FA')
    doc.add_paragraph(); return t
def para(text='',bold=False):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); r=p.add_run(text); r.bold=bold; return p
def nums(items):
    for x in items: doc.add_paragraph(x,style='List Number')
def bullets(items):
    for x in items: doc.add_paragraph(x,style='List Bullet')

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('健康管理平台验收操作与答辩话术手册')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('重复使用运载器健康管理原型系统现场验收准备').italic=True
para('版本 V1.0｜平台地址：http://127.0.0.1:8000/')
para('本手册按当前代码、已接入算法和甲方测试大纲编写。现场演示时只对已有证据作明确承诺；启动时间、实时延迟、连续运行和真实场地联调应按测试大纲现场实测记录。')

doc.add_heading('一 验收总体口径',1)
para('平台形成“数据接入—预处理—异常告警—故障诊断—健康评估—寿命预测—放行决策—故障重演”的闭环，重点覆盖动力、电源和控制典型系统。')
tbl(['项目','推荐口径'],[
['核心功能','已接入并可演示；正式验收仍需现场测试记录。'],
['诊断算法','Hier14、MSFG + TEAMS-RT、PCA–iForest 已接入统一诊断入口。'],
['健康评估','CDPCA-GA、AE-GMM、GCN/RBD 和再飞放行融合路径已提供。'],
['寿命预测','health-main 集成模型和线性 HI 趋势模型已接入。'],
['决策性质','再飞/放行结果是辅助决策，最终由授权人员签署。'],
['未闭环证据','性能实测、真实 MATLAB/Simulink 联调、正式文档和完整测试记录。']])

doc.add_heading('二 验收前准备',1)
nums(['确认 Windows 11、网络和项目目录正常，存在 manage.py、.venv、frontend/dist 和 integrations/algorithm_assets。','使用 start_cmg_platform.bat 启动；手动启动时执行 .venv\\Scripts\\python.exe manage.py runserver 127.0.0.1:8000 --noreload。','打开 http://127.0.0.1:8000/，使用管理员账号登录；另准备一个查看账号验证权限。','确认综合态势页显示 Hier14、健康评估、仿真数据和模型状态。','准备动力故障 CSV、电源故障 CSV 和含 HI_norm 的 RUL 数据集。','准备 MATLAB 采集脚本、接口地址和一次性令牌。','演示前清理无关测试记录，仅保留代表性数据。'])

doc.add_heading('三 推荐演示顺序',1)
tbl(['顺序','页面','操作','结果'],[
['1','综合态势','刷新平台状态','能力链和算法状态'],['2','数据管理','查看、下载、删除数据集','元数据和删除同步'],['3','故障数据集','选择动力/电源故障并生成','CSV、告警、诊断、HI、RUL'],['4','故障诊断','分别运行三种算法','类别、概率、窗口、证据/D矩阵'],['5','健康评估','运行 CDPCA-GA/AE-GMM/GCN-RBD','部件和系统 HI'],['6','寿命预测','选择有效 HI 数据运行 RUL','RUL、趋势、模型来源'],['7','FTA/FMECA','查看逻辑树和测试样例','节点、AND/OR、边、风险记录'],['8','MATLAB 实时采集','创建会话并让 MATLAB 上报','实时曲线、样本和告警'],['9','再飞放行','输入 HI、寿命裕度和告警数','系统 HI、成功概率、阻断项'],['10','故障重演','选择事件并播放','时间窗口曲线和表格']])

doc.add_heading('四 关键操作与话术',1)
doc.add_heading('4.1 数据处理和数据集',2)
para('操作：进入“数据管理”，查看来源、故障类型、注入时间、系统和部件；下载 CSV；管理员删除一条并确认列表立即消失。')
para('话术：“元数据和 CSV 文件绑定保存。删除由后端执行，成功后前端按数据集 ID 立即移除并重新核对，避免数据库和页面不一致。”')
doc.add_heading('4.2 故障注入、告警和诊断',2)
para('操作：进入“实时故障注入与自动诊断”，选择动力/电源/控制、故障模式、注入时刻和严重度，创建任务并播放。')
para('话术：“注入器生成带时间和真值的数据，并按注入时刻触发监测。告警、诊断、隔离、健康评估、RUL 和发布结果属于同一任务上下文，可追溯。”')
para('典型模式：动力泄漏、动力泵效率下降、电源母线欠压、电池容量衰减、控制传感器偏置、控制执行器迟滞。')
doc.add_heading('4.3 三种诊断算法',2)
tbl(['算法','主要作用','输出','边界'],[['Hier14','层次时序诊断','类别、窗口概率和聚合结果','需符合模型字段的 CSV'],['MSFG + TEAMS-RT','D矩阵关联定位','诊断集合、测点证据、候选故障','解释关联，不等同深度分类'],['PCA–iForest','无监督异常筛查','异常窗口、得分、PCA载荷','异常程度不等同具体故障概率']])
para('话术：“三种方法不是混成一个黑盒分数，而是按任务选择：Hier14 偏时序分类，MSFG + TEAMS-RT 偏可解释定位，PCA–iForest 偏异常筛查。”')
doc.add_heading('4.4 健康评估和 RUL',2)
para('健康评估选择数据集和算法后查看 HI 序列及趋势。RUL 预测要求 CSV 含 HI_norm、HI 或 health_index；health-main 集成模型至少需要 51 行 HI。')
para('话术：“健康指数是寿命预测的上游证据。没有有效 HI 的压力、温度原始数据不能直接冒充 RUL 输入。模型单位继承训练数据定义，未标定前不换算成小时或飞行次数。”')
doc.add_heading('4.5 FTA 与 FMECA',2)
para('FTA 展示顶事件、基本事件、AND/OR 逻辑门和边关系；保存后重新加载，验证持久化。FMECA 空库首次进入会生成动力、电源和控制测试样例，编辑后点击“保存全部”。')
para('话术：“FTA 描述故障如何由下层事件组合形成，FMECA 描述故障模式、原因、影响和风险，两者共同支撑诊断解释和维护分析。”')
doc.add_heading('4.6 MATLAB 实时采集',2)
nums(['管理员在“状态监测—遥测数据监测”选择 MATLAB 实时采集。','创建采集会话并复制一次性 Bearer Token 和上报地址。','MATLAB 加入 integrations/matlab 路径，调用 phm_send_telemetry 或 demo_realtime_fault_injection。','页面选择该会话，观察样本数和趋势。','异常告警中心选择同一会话并打开实时检测。','结束采集后导出 CSV，再进入故障诊断。'])
para('话术：“实时采集是会话级 HTTP 接入。平台校验字段、时间递增和数值有效性后持久化，再执行实时规则检测。”')
para('注意：同一会话字段必须一致，time 必须严格递增；远程 MATLAB 将 127.0.0.1 改为平台服务器 IP。')
doc.add_heading('4.7 再飞与放行评估',2)
para('输入动力、电源、测控、导航制导与控制、结构 HI，寿命裕度、严重告警数和放行阈值。')
para('话术：“系统先计算加权健康度和 RBD 等效健康度，再融合量纲适配后的 RF 结果。子系统低于最低阈值、存在严重告警或成功概率低于阈值时暂缓放行。结果是辅助决策，最终由授权人员签署。”')

doc.add_heading('五 甲方常见提问和建议回答',1)
qas=[('数据是真实飞行数据吗？','演示数据区分仿真、离线和实时来源，主要用于联调和回归；甲方真实数据可按相同接口接入。'),('诊断算法用了哪一个？','统一入口支持三种算法，分别承担时序分类、关联定位和异常筛查。'),('准确率是多少？','需用甲方确认的标注集和统一口径现场统计，没有标注集不编造数值。'),('RUL单位是什么？','继承训练数据定义；线性趋势模型单位是样本间隔，未标定不写成小时。'),('MATLAB怎么接入？','管理员创建会话获取地址和令牌，MATLAB通过 phm_send_telemetry 批量上报。'),('实时告警为何没有？','检查是否选择同一会话并打开实时检测，再检查字段名和规则阈值。'),('FTA和FMECA区别？','FTA关注逻辑组合和传播，FMECA关注模式、原因、影响和风险。'),('能保证一秒刷新吗？','系统按批量接收和轮询设计，合同指标需在现场用时间戳实测并记录。'),('是否全部验收通过？','核心功能已完成开发回归；正式验收还需现场性能、真实联调、文档齐套和测试记录签署。')]
for q,a in qas:
    para('甲方：'+q,True); para('建议答：'+a)

doc.add_heading('六 现场测试记录要点',1)
tbl(['测试项','记录内容','通过依据'],[['启动性能','启动脚本到进入工作界面的开始/结束时间，重复3次','每次不超过60秒'],['实时刷新','发送、入库、页面显示时间及样本数','按大纲统计平均和最大延迟'],['实时告警','异常信号、阈值、告警时间和事件','告警、等级、隔离和HI正确'],['诊断接口','三种算法请求和响应 JSON/截图','输入校验和输出字段完整'],['权限','管理员和查看用户操作结果','越权操作被拒绝'],['异常输入','错误密码、空数据、乱序时间、循环边','拒绝并给出提示，不崩溃'],['连续运行','采集时长、样本数、页面切换和导出','无崩溃、无数据错乱']])
para('记录至少包含：用例编号、环境、输入、操作人、开始/结束时间、预期结果、实际结果、判定、问题编号和回归结论。')

doc.add_heading('七 答辩边界',1)
tbl(['不要说','建议改说'],[['准确率已经达到某百分比','需基于甲方标注集按统一口径统计'],['RUL就是多少小时','当前单位继承训练数据，完成标定后换算'],['VAE-GAN完全校准','VAE-GAN作为辅助证据，主判据使用系统健康和量纲适配RF'],['181个用例全部通过','开发回归已通过，现场用例需逐项执行并签署'],['MATLAB可由网页启动','MATLAB独立运行，通过地址和令牌上报']])
doc.add_heading('八 验收结束总结话术',1)
para('“本次演示覆盖数据处理、MATLAB 实时采集、异常告警、三种故障诊断、动力和电源故障注入、结构树、FTA/FMECA、健康评估、RUL、再飞放行和故障重演。平台结果保留数据来源、模型或规则路径及限制说明。启动时间、刷新延迟、连续运行和真实场地联调建议按测试大纲现场记录并签署。”')
doc.add_heading('九 材料打包清单',1)
bullets(['源码、启动脚本和依赖清单；','软件方案、使用手册、数据接口说明、技术研究总结报告；','需求—功能—测试用例追踪表；','现场测试记录、问题闭环表和版本信息；','动力/电源典型故障 CSV、数据字典、FTA/FMECA 和结构树；','Hier14、MSFG + TEAMS-RT、PCA–iForest、健康评估和 RUL 算法说明；','MATLAB 脚本、接口格式和令牌使用说明；','启动时间、实时延迟、连续运行和告警延迟记录。'])
doc.save(OUT); print(OUT)
