import pandas as pd
from docx import Document
from io import BytesIO
import pandas as pd

def parse_table_from_doc(docPath=None):
    report = Document(docPath)
    tables = None
    for tid, table in enumerate(report.tables or []):
        nrows = len(table.rows)
        ncols = len(table.columns)
        res = [[table.cell(i,j).text for j in range(ncols)] for i in range(nrows)]
        if tables is None:
            tables = pd.DataFrame(res[1:], columns=res[0])
        else:
            tables = pd.merge(tables, pd.DataFrame(res[1:], columns=res[0]), how="outer")
    return tables

def read_table(io_, header=0, index_col=None):
    with BytesIO() as io_roaming:
        io_.save(io_roaming)
        io_roaming.seek(0)
        try:
            return pd.read_excel(io_roaming, header=header, index_col=index_col)
        except:
            try:
                io_roaming.seek(0)
                return pd.read_csv(io_roaming, header=header, index_col=index_col, encoding="gbk")
            except:
                try:
                    io_roaming.seek(0)
                    return pd.read_csv(io_roaming, header=header, index_col=index_col, encoding="utf-8")
                except:
                    try:
                        io_roaming.seek(0)
                        table =  parse_table_from_doc(io_roaming)
                        if index_col in table.columns.values:
                            return table.set_index(index_col)
                        elif isinstance(index_col, int) and index_col < len(table.columns.values):
                            index_col = table.columns.values[index_col]
                            return table.set_index(index_col)
                        else:
                            return table
                    except:
                        raise Exception("Unknown Type File")
                        