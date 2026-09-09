import os, time, uuid, re
from io import BytesIO
import numpy as np
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT

if os.path.exists("visio_gui_utils.py"):
    from visio_gui_utils import render_node_pos
else:
    from .visio_gui_utils import render_node_pos
    
TYP_MAP = {"fault-node": [MSO_SHAPE.ROUNDED_RECTANGLE, [1], [3]],
           "test-node": [MSO_SHAPE.OVAL, [2], [6]],
           "OR": [MSO_SHAPE.FLOWCHART_STORED_DATA, [3], [1]],
           "and-node": [MSO_SHAPE.FLOWCHART_DELAY, [1], [3]]}
TYP_QUE = {v[0]: k for k,v in TYP_MAP.items()}

IDKEY = "id"
NAMEKEY = "text"
CLASSKEY = "type"
SRCKEY = "from"
TRGKEY = "to"

def get_uuid(kw=""):
    "生成ID"
    return str(uuid.uuid5(uuid.NAMESPACE_DNS, f"{kw}#%.7f"%time.time()))


def is_auto_shape_type(shape_):
    try:
        if not  "auto_shape_type" in dir(shape_):
            return False
        return not (shape_.auto_shape_type.value is None)
    except:
        return False

def _wash_or_node_pptx(edge, trgRootId, or_source_map, check_map=set()):
    edges = []
    for trgid in or_source_map[trgRootId]:
        if trgid in or_source_map.keys():
            if trgid in check_map:
                continue
            edges.extend(_wash_or_node_pptx(edge, trgid, or_source_map, check_map|{trgRootId, trgid}))
        else:
            edges.append({**edge, TRGKEY: trgid})
    return edges

def faultTree2msfd(struct):
    or_node = [node[IDKEY] for node in struct["nodes"] if node[CLASSKEY].lower() == "or"]
    nodes = [node for node in struct["nodes"] if node[CLASSKEY].lower() != "or"]
    or_source_map = {}
    for edge in struct["edges"]:
        if edge[SRCKEY] in or_node:
            if not edge[SRCKEY] in or_source_map.keys():
                or_source_map[edge[SRCKEY]] = [edge[TRGKEY]]
            else:
                or_source_map[edge[SRCKEY]].append(edge[TRGKEY])
    edges = []
    for edge in struct["edges"]:
        if edge[TRGKEY] in or_source_map.keys():
            edges.extend(_wash_or_node_pptx(edge, edge[TRGKEY], or_source_map, set()))
        elif not edge[SRCKEY] in or_node:
            edges.append(edge)
    return {"nodes": nodes, "edges": edges}

def get_shape_pos(shapes_):
    return [[(shape_.left, shape_.left+shape_.width),
              (shape_.top, shape_.top+shape_.height)] for shape_ in shapes_
             if (not "begin_x" in dir(shape_)) and \
                 is_auto_shape_type(shape_) and shape_.auto_shape_type.value in TYP_QUE.keys()]

def get_connector_rule(edge, shape_pos):
    edge_end = [edge.end_x, edge.end_y]
    end_id = np.argmin([max(0, shape_p[0][0]-1 - edge_end[0], edge_end[0] - shape_p[0][1]+1)**2 +\
              max(0, shape_p[1][0]-1 - edge_end[1], edge_end[1] - shape_p[1][1]+1)**2  for shape_p in shape_pos])
    edge_start = [edge.begin_x, edge.begin_y]
    start_id = np.argmin([max(0, shape_p[0][0]-1 - edge_start[0], edge_start[0] - shape_p[0][1]+1)**2 +\
              max(0, shape_p[1][0]-1 - edge_start[1], edge_start[1] - shape_p[1][1]+1)**2  for shape_p in shape_pos])
    return int(start_id), int(end_id)
    
def parse_pptx(path):
    """从visio对应的path/ByteIO中读取单幅流图"""
    prs = Presentation(path)
    result = {"nodes": [], "edges": []}
    shapes_ =  prs.slides[0].shapes
    shape_pos_ = get_shape_pos(shapes_)
    shape_ids_ = [f"SHAPE#{shape_id}#{shape_.shape_id}" for shape_id, shape_ in enumerate(shapes_)
                 if (not "begin_x" in dir(shape_)) and \
                 is_auto_shape_type(shape_) and shape_.auto_shape_type.value in TYP_QUE.keys()]
    for nodeId, node in enumerate(shapes_):
        if (not "begin_x" in dir(node)) and is_auto_shape_type(node) and node.auto_shape_type.value in TYP_QUE.keys():
            type_ =  TYP_QUE[node.auto_shape_type.value]
            result["nodes"].append({
                NAMEKEY: node.text,
                CLASSKEY: type_,
                IDKEY: f"SHAPE#{nodeId}#{node.shape_id}",
            } ) 
        elif "begin_x" in dir(node):
            start_id, end_id = get_connector_rule(node, shape_pos_)
            result["edges"].append({
                "from_id": shape_ids_[start_id],
                "to_id": shape_ids_[end_id],
                })
    nodeMap = {node["id"]: {
        IDKEY: get_uuid(f"{node['text']}#{nodeId}"),
        NAMEKEY: node["text"],
        CLASSKEY: node["type"]
        } for nodeId, node in enumerate(result["nodes"])}
    return {"nodes": list(nodeMap.values()),
            "edges": [{SRCKEY: nodeMap[edge["from_id"]][IDKEY],
                       "from_text": nodeMap[edge["from_id"]][NAMEKEY],
                       "to_text": nodeMap[edge["to_id"]][NAMEKEY],
                       TRGKEY: nodeMap[edge["to_id"]][IDKEY]} for edge in result["edges"]
                      if edge["from_id"] in nodeMap.keys() and edge["to_id"] in nodeMap.keys()]}
    
def msfd2faultTree(struct):
    edgeMap = {}
    nodes = [{**node, NAMEKEY: node[NAMEKEY] if isinstance(node[NAMEKEY], str) else node[NAMEKEY]["value"]} for node in struct["nodes"]]
    edges_ = [{
        SRCKEY: nodes[k][IDKEY], 
        TRGKEY: nodes[v_][IDKEY]
        } for k, v in enumerate(struct["edges"]) for v_ in v]
    edges = []
    and_node_keys = [node[IDKEY] for node in struct["nodes"] if node[CLASSKEY].lower() == "and-node"]
    for edge in edges_:
        start_id = edge[SRCKEY]
        end_id = edge[TRGKEY]
        if end_id in and_node_keys:
            edges.append(edge)
            continue
        if not end_id in edgeMap.keys():
            edgeMap[end_id] = []
        edgeMap[end_id].append(start_id)
    cnt = 0
    for k, vs in edgeMap.items():
        new_uuid = get_uuid(f"OR#{cnt}#{time.time()}".replace(".", "_"))
        nodes.append({NAMEKEY: "", CLASSKEY: "OR", IDKEY: new_uuid,})
        edges.append({SRCKEY: new_uuid,
                      TRGKEY: k})
        for v in set(vs):
            edges.append({SRCKEY: v,
                          TRGKEY: new_uuid})
        cnt += 1
    return {"nodes": nodes, "edges": edges}

def _get_bytes_length(text):
    text_ = re.sub(r"[0-9a-zA-Z@\.\+\-\_]", "", text)
    length = len(text) - len(text_)
    return max(0, length*0.83 + len(text_))+1.5

class pptxGramManage:
    def __init__(self):
        self.re_init()

    def re_init(self):
        self._prs = Presentation()
        slide_layout = self._prs.slide_layouts[1]
        self._slide = self._prs.slides.add_slide(slide_layout)
        self._cnt = len(self._slide.shapes)
        self._node_map = {}
        self._start_map = {}
        self._end_map = {}
        self._pos_map = {}

    def add_node(self, pos, typ, text, ref=None, px=0.4):
        left, top = pos
        width = Inches(px*_get_bytes_length(text))
        height = Inches(px)
        shape_x, in_conn, out_conn = TYP_MAP[typ]
        shape = self._slide.shapes.add_shape(shape_x,
                                             Inches(left), Inches(top),
                                             width, height)
        ref = ref or f"{self._cnt}#{text}"
        self._node_map[ref] = self._cnt
        self._start_map[ref] = in_conn
        self._end_map[ref] = out_conn
        self._pos_map[ref] = [pos[0]+px*_get_bytes_length(text)/2, pos[1]+px/2]
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255,255,255)
        shape.text = text
        shape.text_frame.paragraphs[0].alignment = PP_PARAGRAPH_ALIGNMENT.CENTER
        shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0,0,0)
        shape.line.color.rgb = RGBColor(0,0,0)
        if typ == "OR":
            shape.rotation = 180
        self._cnt += 1

    def add_edge(self, src_ref, trg_ref):
        shape = self._slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT,
                                                 Inches(0), Inches(0),
                                                 Inches(0), Inches(0))
        shape.line.end_arrowhead_length = shape.end_arrowhead_width = Inches(1)
        shape.line.end_arrowhead_style = 'Triangle'
        src_x, _ = self._pos_map[src_ref]
        trg_x, _ = self._pos_map[trg_ref]
        start_choice = self._start_map[src_ref]
        end_choice = self._end_map[trg_ref]
        if src_x > trg_x:
            shape._connect_begin_to(self._slide.shapes[self._node_map[src_ref]], start_choice[0])
            shape._connect_end_to(self._slide.shapes[self._node_map[trg_ref]], end_choice[-1])
        else:
            shape._connect_begin_to(self._slide.shapes[self._node_map[src_ref]], start_choice[-1])
            shape._connect_end_to(self._slide.shapes[self._node_map[trg_ref]], end_choice[0])
        shape._move_begin_to_cxn(self._slide.shapes[self._node_map[src_ref]], 1)
        shape._move_end_to_cxn(self._slide.shapes[self._node_map[trg_ref]], 3)
        shape.line.color.rgb = RGBColor(0,0,0)
        self._cnt += 1
        pass
        
    def save(self, path):
        self._prs.save(path)

def drawStruct(struct, path,
               pixelSize=0.4, width=3, height=1.5, padding=0.2):
    struct = msfd2faultTree(struct)
    nodes = struct.get("nodes", []); edges = struct.get("edges", [])
    nodes = render_node_pos(nodes, edges,
                            width=width, height=height, padding=padding,
                            srcKey=SRCKEY, trgKey=TRGKEY, idKey=IDKEY)
    v = pptxGramManage()
    for node in nodes:
        v.add_node([node['x'],node['y']],
                   node[CLASSKEY],node[NAMEKEY],node[IDKEY],
                   px=pixelSize)
    for edge in edges:
        v.add_edge(edge[SRCKEY],edge[TRGKEY]) 
    v.save(path)

if __name__ == "__main__":
    res = parse_pptx("pic.pptx")
    res_msfd = faultTree2msfd(res)
    drawStruct(res_msfd, "1.pptx")
